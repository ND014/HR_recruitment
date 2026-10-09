import docx
import io
from typing import Union, BinaryIO

def extract_text_from_docx(file_source: Union[str, bytes, BinaryIO]) -> str:
    """
    Extracts text from a DOCX file, preserving paragraphs and table contents.
    
    Args:
        file_source: File path, bytes, or file-like object.
    
    Returns:
        Extracted plain text string.
    """
    if isinstance(file_source, bytes):
        doc = docx.Document(io.BytesIO(file_source))
    elif hasattr(file_source, "read"):
        file_source.seek(0)
        doc = docx.Document(file_source)
    else:
        doc = docx.Document(str(file_source))

    extracted_lines = []

    # Extract paragraphs
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            extracted_lines.append(text)

    # Extract tables (often used for education, skills, experience)
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                # Deduplicate identical adjacent cells (merged cells)
                deduped = []
                for cell in row_cells:
                    if not deduped or deduped[-1] != cell:
                        deduped.append(cell)
                extracted_lines.append(" | ".join(deduped))

    return "\n".join(extracted_lines)
