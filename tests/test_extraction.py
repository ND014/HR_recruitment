import os
import sys
from PIL import Image, ImageDraw, ImageFont
import docx

# Add project root to sys.path
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE_DIR)

from parsers.docx_parser import extract_text_from_docx
from parsers.image_parser import extract_text_from_image
from parsers.pdf_parser import extract_text_from_pdf
from extractor.nlp_pipeline import parse_resume

def create_sample_docx(filepath: str):
    doc = docx.Document()
    doc.add_heading("Alex Morgan", level=0)
    doc.add_paragraph("alex.morgan@example.com | +1 (555) 234-5678 | Age: 29")
    
    doc.add_heading("Education", level=1)
    doc.add_paragraph("B.Tech in Computer Science\nStanford University, 2015 - 2019\nGPA: 3.8 / 4.0")
    doc.add_paragraph("Full Stack Web Development Certificate\nSchool of Code Online Bootcamp, 2020")
    
    doc.add_heading("Work Experience", level=1)
    doc.add_paragraph("Senior Software Engineer at Acme Corp\nJan 2021 - Present")
    doc.add_paragraph("- Designed and deployed microservices architecture handling 10M requests daily.")
    doc.add_paragraph("- Mentored 4 junior engineers.")
    
    doc.add_paragraph("Software Engineer at Beta Solutions\nJun 2019 - Dec 2020")
    doc.add_paragraph("- Built scalable RESTful APIs with Python and FastAPI.")
    
    doc.save(filepath)

def create_sample_image(filepath: str):
    # Create white canvas
    img = Image.new("RGB", (800, 600), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    
    # Simple default font
    resume_text = (
        "SARAH CONNOR\n"
        "sarah.c@sky.net | +91 9876543210 | Age: 31\n\n"
        "EDUCATION\n"
        "Master of Science in Data Science\n"
        "MIT University, 2016 - 2018\n\n"
        "WORK EXPERIENCE\n"
        "Lead Data Scientist at Cyberdyne Systems\n"
        "2019 - Present\n"
        "- Developed anomaly detection algorithms.\n"
    )
    
    draw.text((40, 40), resume_text, fill=(0, 0, 0))
    img.save(filepath)

def test_docx_pipeline():
    docx_path = os.path.join(BASE_DIR, "tests", "sample_alex.docx")
    create_sample_docx(docx_path)
    
    raw_text = extract_text_from_docx(docx_path)
    assert len(raw_text) > 0, "DOCX extraction produced empty text"
    
    result = parse_resume(raw_text)
    print("DOCX Extraction Result:")
    print("  Name:", result["name"])
    print("  Age:", result["age"])
    print("  Phone:", result["phone_number"])
    print("  Education:", result["education"])
    print("  Work Experience:", result["work_experience"])
    
    assert "Alex Morgan" in result["name"] or "Alex" in result["name"], f"Name failed: {result['name']}"
    assert "29" in result["age"], f"Age failed: {result['age']}"
    assert "555" in result["phone_number"], f"Phone failed: {result['phone_number']}"
    assert len(result["education"]) > 0, "Education failed"
    assert len(result["work_experience"]["positions"]) > 0, "Work experience failed"
    print("[OK] DOCX Test PASSED\n")

def test_image_pipeline():
    img_path = os.path.join(BASE_DIR, "tests", "sample_sarah.png")
    create_sample_image(img_path)
    
    raw_text = extract_text_from_image(img_path)
    assert len(raw_text) > 0, "Image OCR produced empty text"
    
    result = parse_resume(raw_text)
    print("Image OCR Extraction Result:")
    print("  Name:", result["name"])
    print("  Age:", result["age"])
    print("  Phone:", result["phone_number"])
    print("  Education:", result["education"])
    print("  Work Experience:", result["work_experience"])
    
    assert "31" in result["age"], f"Age failed: {result['age']}"
    assert "9876543210" in result["phone_number"], f"Phone failed: {result['phone_number']}"
    print("[OK] Image OCR Test PASSED\n")

def test_pdf_pipeline():
    # Test PDF parsing using image-based PDF created with PIL
    pdf_path = os.path.join(BASE_DIR, "tests", "sample_scanned.pdf")
    img = Image.new("RGB", (800, 600), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    resume_text = (
        "DAVID MILLER\n"
        "david.m@example.com | +1 800-555-0199 | Age: 34\n\n"
        "EDUCATION\n"
        "Bachelor of Science in Electrical Engineering\n"
        "University of California, Berkeley - 2010 - 2014\n\n"
        "EXPERIENCE\n"
        "DevOps Engineer at Cloud Solutions\n"
        "2018 - Present\n"
        "- Kubernetes infrastructure deployment.\n"
    )
    draw.text((40, 40), resume_text, fill=(0, 0, 0))
    img.save(pdf_path, "PDF", resolution=100.0)
    
    raw_text = extract_text_from_pdf(pdf_path)
    assert len(raw_text) > 0, "PDF extraction produced empty text"
    
    result = parse_resume(raw_text)
    print("PDF Extraction Result:")
    print("  Name:", result["name"])
    print("  Age:", result["age"])
    print("  Phone:", result["phone_number"])
    print("  Education:", result["education"])
    print("  Work Experience:", result["work_experience"])
    
    assert "34" in result["age"], f"Age failed: {result['age']}"
    assert "555" in result["phone_number"], f"Phone failed: {result['phone_number']}"
    print("[OK] PDF Test PASSED\n")

def test_abbreviations_and_synonyms():
    raw_sample = """
    PRIYA SHARMA
    priya.sharma@example.com | +91 9123456789 | Age: 27

    CAREER SUMMARY
    SDE-2 at Amazon Web Services
    Feb 2022 - Present
    - Developed high-throughput distributed caching systems.
    - Reduced latency by 35%.

    SWE Intern at Google
    May 2021 - Aug 2021
    - Implemented search ranking algorithms.

    ACADEMIC QUALIFICATIONS
    MBA in Finance & Technology
    Indian Institute of Management, Bangalore (2020 - 2022)
    CGPA: 3.9 / 4.0

    BSc Computer Science
    King's College London (2016 - 2019)
    First Class Honours

    BA in Economics
    Delhi University (2013 - 2016)

    Fullstack Web Developer Certificate
    Udemy Online Bootcamp (2020)
    """
    
    result = parse_resume(raw_sample)
    print("Abbreviations & Synonyms Test Result:")
    print("  Name:", result["name"])
    print("  Age:", result["age"])
    print("  Phone:", result["phone_number"])
    print("  Education:", result["education"])
    print("  Work Experience:", result["work_experience"])
    
    # Assertions
    assert "Priya Sharma" in result["name"], f"Name failed: {result['name']}"
    assert "27" in result["age"], f"Age failed: {result['age']}"
    assert "9123456789" in result["phone_number"], f"Phone failed: {result['phone_number']}"
    
    # Education assertions
    edu_insts = [e["institution"] for e in result["education"]]
    edu_degs = [e["degree"] for e in result["education"]]
    print("  Extracted Edu Institutions:", edu_insts)
    print("  Extracted Edu Degrees:", edu_degs)
    
    assert any("King" in inst for inst in edu_insts), "King's College missing"
    assert any("Delhi University" in inst for inst in edu_insts), "Delhi University missing"
    assert not any("Udemy" in inst for inst in edu_insts), "Udemy should be excluded"
    assert any("BSc" in deg for deg in edu_degs), "BSc missing"
    assert any("BA" in deg for deg in edu_degs), "BA missing"
    
    # Work Experience assertions
    roles = [p["role"] for p in result["work_experience"]["positions"]]
    companies = [p["company"] for p in result["work_experience"]["positions"]]
    print("  Extracted Roles:", roles)
    print("  Extracted Companies:", companies)
    assert any("SDE-2" in r or "Sde" in r for r in roles), "SDE-2 role missing"
    assert any("SWE Intern" in r or "Swe" in r for r in roles), "SWE Intern role missing"
    assert any("Amazon" in c for c in companies), "Amazon company missing"
    assert any("Google" in c for c in companies), "Google company missing"
    
    print("[OK] Abbreviations & Synonyms Test PASSED\n")

if __name__ == "__main__":
    print("Running automated extraction tests...")
    test_docx_pipeline()
    test_image_pipeline()
    test_pdf_pipeline()
    test_abbreviations_and_synonyms()
    print("All extraction tests successfully passed!")
